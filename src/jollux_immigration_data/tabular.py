"""Small, dependency-light readers for disclosure tables."""

from __future__ import annotations

import csv
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import openpyxl


def clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = " ".join(value.split())
        return value or None
    return value


def _normalize_header(value: Any) -> str:
    return str(value or "").strip().upper()


def iter_rows(path: str | Path) -> Iterator[dict[str, Any]]:
    """Yield normalized-header row dictionaries from XLSX or CSV input."""
    source = Path(path)
    suffix = source.suffix.casefold()
    if suffix == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if reader.fieldnames is None:
                raise ValueError("Input has no header row")
            headers = [_normalize_header(name) for name in reader.fieldnames]
            for values in reader:
                yield {header: clean(values.get(original)) for original, header in zip(reader.fieldnames, headers)}
        return
    if suffix != ".xlsx":
        raise ValueError("Input must be an .xlsx or .csv file")

    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        try:
            headers = [_normalize_header(value) for value in next(rows)]
        except StopIteration as error:
            raise ValueError("Input has no header row") from error
        for values in rows:
            yield {header: clean(values[index] if index < len(values) else None) for index, header in enumerate(headers)}
    finally:
        workbook.close()


def require_columns(first_row: dict[str, Any], required: set[str]) -> None:
    missing = sorted(required - first_row.keys())
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")
